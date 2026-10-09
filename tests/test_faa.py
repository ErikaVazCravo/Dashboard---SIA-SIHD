import builtins
from io import BytesIO
from pathlib import Path
import unittest
from unittest.mock import patch
import pandas as pd
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest
from processamento import excel
from processadores.faa import processar_faa


def relatorio(unidade, especialidades, quantidades, cabecalho=3):
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine='xlsxwriter', engine_kwargs={'options': {'in_memory': True}}) as writer:
        pd.DataFrame({'Unidade Prestadora': [unidade] * len(especialidades),
            'Especialidade': especialidades, 'Quantidade': quantidades}).to_excel(
                writer, sheet_name='Relatório', index=False, startrow=cabecalho)
        writer.sheets['Relatório'].write(0, 0, 'Relatório FAA')
    return buffer.getvalue()


class FaaTest(unittest.TestCase):
    def test_contagem_abas_e_agrupamentos_sem_gravar(self):
        arquivos = {
            '2064731': [('central.xlsx', relatorio('UPA Central',
                ['AUXILIAR DE ENFERMAGEM', 'ENFERMEIRA (O)', 'PEDIATRIA'], [3, 7, 2]))],
            '4047184': [('xavier.xlsx', relatorio('UPA Vila Xavier', ['FARMACEUTICO'], [5]))],
            '9099581': [('selmi.xlsx', relatorio('UPA Selmi Dei', ['MEDICO GENERALISTA'], [4]))],
        }
        abrir = builtins.open
        def somente_leitura(arquivo, mode='r', *args, **kwargs):
            if any(letra in mode for letra in 'wax+'):
                raise AssertionError('Tentativa de gravar dados no disco.')
            return abrir(arquivo, mode, *args, **kwargs)
        with patch('builtins.open', side_effect=somente_leitura), patch('io.open', side_effect=somente_leitura):
            dados, resumo = processar_faa('08/2026', arquivos)
        self.assertEqual(resumo['Qtde.'].sum(), 21)
        wb = load_workbook(BytesIO(dados))
        self.assertEqual(wb.sheetnames, ['ESPECIALIDADES_PADRONIZADAS', 'ORIGINAL'])
        linhas = list(wb['ESPECIALIDADES_PADRONIZADAS'].values)
        enfermagem = next(row for row in linhas[1:] if row[5] == 'ENFERMAGEM')
        self.assertEqual(enfermagem, (2026, 8, '08/2026', '2064731', 'UPA Central', 'ENFERMAGEM', 10))
        self.assertEqual(wb['ORIGINAL'].max_row, 6)
        self.assertEqual(sum(row[6] for row in list(wb['ORIGINAL'].values)[1:]), 21)

    def test_apenas_uma_upa_e_mais_de_um_arquivo(self):
        dados, resumo = processar_faa('07/2026', {'2064731': [
            ('a.xlsx', relatorio('UPA Central', ['PEDIATRIA'], [2])),
            ('b.xlsx', relatorio('UPA Central', ['PEDIATRIA'], [3]))]})
        self.assertEqual(len(resumo), 1)
        self.assertEqual(resumo['Qtde.'].iloc[0], 5)
        self.assertEqual(load_workbook(BytesIO(dados))['ORIGINAL']['G2'].value, 5)

    def test_rejeita_duplicado(self):
        dados = relatorio('UPA Central', ['PEDIATRIA'], [2])
        with self.assertRaisesRegex(ValueError, 'mais de uma vez'):
            processar_faa('08/2026', {'2064731': [('a.xlsx', dados), ('b.xlsx', dados)]})

    def test_rejeita_upa_no_campo_errado(self):
        with self.assertRaisesRegex(ValueError, 'UPA diferente'):
            processar_faa('08/2026', {'2064731': [('a.xlsx', relatorio('UPA Selmi Dei', ['PEDIATRIA'], [2]))]})

    def test_rejeita_quantidades_invalidas(self):
        for valor in ['abc', -1, 1.5]:
            with self.subTest(valor=valor), self.assertRaisesRegex(ValueError, 'inteiros não negativos'):
                processar_faa('08/2026', {'2064731': [('a.xlsx', relatorio('UPA Central', ['PEDIATRIA'], [valor]))]})

    def test_rejeita_cabecalho_incorreto(self):
        with self.assertRaisesRegex(ValueError, 'Cabeçalho'):
            processar_faa('08/2026', {'2064731': [('a.xlsx', excel({'Teste': pd.DataFrame({'Nome': ['ABC']})}))]})

    def test_sem_arquivos(self):
        with self.assertRaisesRegex(ValueError, 'pelo menos um'):
            processar_faa('08/2026', {})

    def test_interface_faa_independente(self):
        app = AppTest.from_file(str(Path(__file__).resolve().parents[1] / 'app.py')).run()
        app.sidebar.radio[0].set_value('UPAs · FAA').run()
        self.assertFalse(app.exception)
        self.assertEqual(len(app.get('file_uploader')), 3)
        self.assertEqual(len(app.text_input), 1)
        self.assertTrue(app.button[1].disabled)
        app.session_state['resultados'] = {'sia.xlsx': b'original'}
        app.session_state['faa_resultado'] = b'faa'
        app.session_state['faa_resumo'] = pd.DataFrame({'Prestador': ['UPA Central'], 'Qtde.': [10]})
        app.run()
        self.assertEqual(len(app.get('download_button')), 1)
        app.button[0].click().run()
        self.assertFalse(app.exception)
        with self.assertRaises(KeyError):
            app.session_state['faa_resultado']
        self.assertEqual(app.session_state['resultados'], {'sia.xlsx': b'original'})
        app.sidebar.radio[0].set_value('SIA + SIHD').run()
        self.assertFalse(app.exception)
        self.assertEqual(len(app.get('file_uploader')), 11)
        with self.assertRaises(KeyError):
            app.session_state['resultados']
