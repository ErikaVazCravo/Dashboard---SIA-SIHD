from io import BytesIO
from pathlib import Path
from uuid import uuid4
import builtins
import unittest
from unittest.mock import patch
from zipfile import ZipFile
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest
from processamento import processar, compactar, validar_competencia, internacoes, excel
import pandas as pd


def largura(campos, tamanho=110):
    linha = list(' ' * tamanho)
    for inicio, texto in campos.items():
        linha[inicio:inicio + len(texto)] = texto
    return ''.join(linha)


def pdf(texto):
    """PDF textual mínimo em memória; dispensa bibliotecas de criação de PDF."""
    instrucoes = ['BT /F1 9 Tf 30 800 Td 12 TL']
    for linha in texto.splitlines():
        linha = linha.replace('\\', '\\\\').replace('(', '\\(').replace(')', '\\)')
        instrucoes.append(f'({linha}) Tj T*')
    conteudo = ('\n'.join(instrucoes) + '\nET').encode('cp1252')
    objetos = [b'<< /Type /Catalog /Pages 2 0 R >>',
        b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
        b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 900 850] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>',
        b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>',
        f'<< /Length {len(conteudo)} >>\nstream\n'.encode() + conteudo + b'\nendstream']
    dados = b'%PDF-1.4\n'
    offsets = [0]
    for i, obj in enumerate(objetos, 1):
        offsets.append(len(dados))
        dados += f'{i} 0 obj\n'.encode() + obj + b'\nendobj\n'
    inicio = len(dados)
    dados += b'xref\n0 6\n0000000000 65535 f \n'
    dados += b''.join(f'{offset:010} 00000 n \n'.encode() for offset in offsets[1:])
    dados += f'trailer << /Size 6 /Root 1 0 R >>\nstartxref\n{inicio}\n%%EOF'.encode()
    return dados


def dados_teste(pasta):
    principal, secundario = '0901010010', '0301010010'
    linhas = []
    for codigo, nome, valor in [(principal, 'OCI TESTE', '000000010000'), (secundario, 'CONSULTA TESTE', '000000002000')]:
        linhas.append(largura({0: codigo, 10: nome, 294: valor, 330: '202609'}, 336))
    (pasta / 'tb_procedimento.txt').write_text('\n'.join(linhas), encoding='latin-1')
    (pasta / 'rl_procedimento_compativel.txt').write_text(principal+'06'+secundario+'06'+'1'+'0001'+'202609', encoding='latin-1')
    unidade = 'Unidade 2082527 SANTA CASA'
    sint_q = unidade + '\n' + largura({0: 'TOTAL DA UNIDADE', 75: '3', 88: '2'})
    sint_v = unidade + '\n' + largura({0: 'TOTAL DA UNIDADE', 75: '300,00', 88: '200,00'})
    anal_q = unidade + '\n' + largura({0: '202609', 7: principal, 39: '3', 51: '2'})
    anal_v = unidade + '\n' + largura({0: '202609', 7: principal, 39: '300,00', 54: '200,00'})
    apac = '\n'.join(['Estab.Exec. : 2082527 - SANTA CASA', 'APAC: 352667960493-2',
        'Processamento:09/2026', 'Producao:09/2026',
        'SQ PROC. CBO Qt.Prz. Equipe Cnes Terc. Vl.Prz. Qt.Apvd Vl.Apvd SITUACAO',
        largura({1:'01', 4:principal, 26:'1', 56:'100,00', 66:'1', 73:'100,00'}, 103),
        largura({1:'02', 4:secundario, 26:'1', 44:'2082527', 56:'20,00', 66:'1', 73:'20,00'}, 103), ' ' * 30])
    resumo = '\n'.join(['RESUMO DOS VALORES APROVADOS POR CNES', 'Competência: 09/2026',
        '2082527 - SANTA CASA 3 2 0,00',
        'Valor sem complemento 100,00 0,00 0,00 0,00 100,00',
        'Complemento federal 0,00 0,00 0,00 0,00 0,00',
        'Complemento local 0,00 0,00 0,00 0,00 0,00',
        'Total 80,00 20,00 100,00', 'Total do Gestor', '3 2 0,00 80,00 20,00 100,00'])
    def procedencia(qtd, valor):
        return pdf('\n'.join(['DEMONSTRATIVO DE PROCEDÊNCIA', 'Competência: 09/2026',
            'CNES 2082527', 'MUN. ORIGEM', f'350320 ARARAQUARA Cirúrgico {qtd} {valor} {valor}', 'Total do CNES']))
    analitico = '\n'.join(['DEMONSTRATIVO DE PROCEDÊNCIA ANALÍTICO', 'Competência: 09/2026',
        f'350320 ARARAQUARA / SP 3526000000001 {secundario} - CONSULTA TESTE MÉDIA FAEC 2082527 - SANTA CASA 20,00',
        'QTD TOTAL DE AIH 1', 'TOTAL DO MUNICIPIO DE ORIGEM 20,00',
        'TOTAL DO MUNICIPIO EXECUTANTE 1 20,00 0,00 0,00 0,00 20,00'])
    return dict(rsproced=sint_v.encode(), rsprocqt=sint_q.encode(), raproced=anal_v.encode(),
        raprocqt=anal_q.encode(), r0570=apac.encode(), resumo_sihd=pdf(resumo),
        sem_filtro=procedencia(6, '600,00'), eletivo=procedencia(2, '200,00'),
        urgencia=procedencia(3, '300,00'), pmae=pdf(analitico))


class ProcessamentoTest(unittest.TestCase):
    def setUp(self):
        self.pasta = Path(__file__).resolve().parent / ('fixture_' + uuid4().hex)
        self.pasta.mkdir()
        def limpar_fixture():
            for arquivo in self.pasta.iterdir():
                arquivo.unlink()
            self.pasta.rmdir()
        self.addCleanup(limpar_fixture)
        self.arquivos = dados_teste(self.pasta)

    def test_pipeline_sem_gravar_arquivos(self):
        antes = {p.name: p.read_bytes() for p in self.pasta.iterdir()}
        abrir = builtins.open
        def somente_leitura(arquivo, mode='r', *args, **kwargs):
            if any(c in mode for c in 'wax+'):
                raise AssertionError(f'Tentativa de escrita em disco: {arquivo}')
            return abrir(arquivo, mode, *args, **kwargs)
        with patch('builtins.open', side_effect=somente_leitura), patch('io.open', side_effect=somente_leitura):
            saidas, _ = processar(str(self.pasta), '09/2026', self.arquivos)
        self.assertEqual(len(saidas), 8)
        self.assertEqual(antes, {p.name: p.read_bytes() for p in self.pasta.iterdir()})
        for nome, dados in saidas.items():
            wb = load_workbook(BytesIO(dados))
            self.assertGreater(wb.active.max_row, 1, nome)
        oci = load_workbook(BytesIO(saidas['OCI_Producao_Prestadores.xlsx']))
        self.assertEqual(oci.sheetnames, ['DADOS_OCI', 'Vl_SIGTAP', 'Vl_APAC_SIGTAP', 'Total_Terceiros'])
        intern = load_workbook(BytesIO(saidas['SIHD_Internacoes.xlsx']))
        linhas = list(intern['Total Municipios'].values)
        cabecalho = linhas[0]
        outros = next(l for l in linhas[1:] if l[cabecalho.index('Tipo de Atendimento')] == 'Outros Tipos')
        self.assertEqual(outros[cabecalho.index('Qtde. de AIHs')], 1)
        self.assertEqual(outros[cabecalho.index('Valor')], 100)
        aih = load_workbook(BytesIO(saidas['SIHD_AIHs_Procedimentos.xlsx']))
        self.assertEqual(aih['AIHs']['I2'].value, '0301010010')
        with ZipFile(BytesIO(compactar(saidas))) as zipfile:
            self.assertEqual(set(zipfile.namelist()), set(saidas))

    def test_competencia_incorreta(self):
        with self.assertRaisesRegex(ValueError, 'competência'):
            processar(str(self.pasta), '08/2026', self.arquivos)

    def test_diferenca_negativa(self):
        self.arquivos['sem_filtro'] = self.arquivos['eletivo']
        with self.assertRaisesRegex(ValueError, 'excede'):
            internacoes(self.arquivos, '09/2026')

    def test_faltante(self):
        self.arquivos.pop('r0570')
        with self.assertRaisesRegex(ValueError, 'R0570'):
            processar(str(self.pasta), '09/2026', self.arquivos)

    def test_producao_sia_anterior_ao_processamento(self):
        for key in ('raproced', 'raprocqt'):
            self.arquivos[key] = self.arquivos[key].replace(b'202609', b'202608')
        saidas, _ = processar(str(self.pasta), '09/2026', self.arquivos)
        ws = load_workbook(BytesIO(saidas['GR_Procedimentos_SIA.xlsx'])).active
        self.assertEqual(ws['A2'].value, '09/2026')
        self.assertEqual(ws['D2'].value, '08/2026')

    def test_texto_excel_permanece_texto(self):
        ws = load_workbook(BytesIO(excel({'Teste': pd.DataFrame({'Nome': ['=1+1'], 'CNES': ['0012345']})}))).active
        self.assertEqual(ws['A2'].data_type, 's')
        self.assertEqual(ws['B2'].value, '0012345')

    def test_relatorio_sia_nao_reconhecido(self):
        self.arquivos['rsproced'] = b'arquivo incorreto'
        with self.assertRaisesRegex(ValueError, 'layout'):
            processar(str(self.pasta), '09/2026', self.arquivos)

    def test_formato_competencia(self):
        for comp in ['13/2026', '9/2026', '', '202609']:
            with self.assertRaises(ValueError):
                validar_competencia(comp)

    def test_interface(self):
        app = AppTest.from_file(str(Path(__file__).resolve().parents[1] / 'app.py')).run()
        self.assertFalse(app.exception)
        self.assertEqual(len(app.get('file_uploader')), 11)
        self.assertTrue(app.button[1].disabled)
        app.session_state['resultados'] = {'exemplo.xlsx': b'teste'}
        app.button[0].click().run()
        self.assertFalse(app.exception)
        with self.assertRaises(KeyError):
            app.session_state['resultados']


if __name__ == '__main__':
    unittest.main()
