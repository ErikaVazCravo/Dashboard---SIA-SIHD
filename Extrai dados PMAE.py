"""Extrai AIHs do DEMONSTRATIVO DE PROCEDÊNCIA ANALÍTICO do SIHD.

Sem argumentos, abre a seleção de PDFs e do Excel de saída.
Terminal: python "Extrai dados PMAE.py" entrada.pdf --saida SIHD_AIHs_Procedimentos.xlsx
Dependências: pdfplumber, pandas e openpyxl.
"""
import argparse
from decimal import Decimal
from pathlib import Path
import re
import unicodedata

import pandas as pd
import pdfplumber


PRESTADORES = {
    "2079763": "CAIRBAR SCHUTEL",
    "6943284": "MATERNIDADE GOTA DE LEITE",
    "2082527": "SANTA CASA",
    "9267263": "UNIDADE DE RETAGUARDA",
}
ORIGEM = "Municipio / UF de ORIGEM"
COD_MUNICIPIO = "Cód. Município"
COLUNAS = ["ANO", "MÊS", "Comp", "CNES", "Prestador", COD_MUNICIPIO, ORIGEM, "AIH",
           "Cod. Procedimento", "Procedimento", "Valor"]
VALOR = r"-?(?:\d{1,3}(?:\.\d{3})+|\d+),\d{2}"
LINHA = re.compile(
    rf"^(?P<municipio>\d{{6}})\s+(?P<origem>.+?\s*/\s*[A-Z]{{2}})\s+"
    rf"(?P<aih>\d{{13}})\s+(?P<codigo>\d{{10}})\s*-\s*"
    rf"(?P<descricao>.+?)\s+(?:ALTA|M[EÉ]DIA|BAIXA)\s+(?:FAEC|MAC)\s+"
    rf"(?P<cnes>\d{{7}})\s*-\s*(?P<prestador>.+?)\s+(?P<valor>{VALOR})$",
    re.IGNORECASE,
)


def normalizar(texto):
    return ''.join(c for c in unicodedata.normalize('NFKD', texto)
                   if not unicodedata.combining(c)).upper()


def converter_valor(valor):
    return Decimal(valor.replace('.', '').replace(',', '.'))


def extrair_texto_pdf(caminho_pdf):
    linhas = []
    with pdfplumber.open(caminho_pdf) as pdf:
        for numero, pagina in enumerate(pdf.pages, 1):
            texto = pagina.extract_text()
            if not texto or not texto.strip():
                raise ValueError(f"Página {numero} sem texto extraível. Verifique se o PDF precisa de OCR.")
            linhas.extend(' '.join(l.split()) for l in texto.splitlines())
    return linhas


def extrair_competencia(linha):
    encontrado = re.search(r'COMPETENCIA\s*:\s*(\d{2}/\d{4})', normalizar(linha))
    return encontrado.group(1) if encontrado else ''


def interpretar_linha_producao(linha):
    encontrado = LINHA.fullmatch(' '.join(linha.split()))
    if not encontrado:
        return None
    d = encontrado.groupdict()
    return {
        'CNES': d['cnes'],
        'Prestador': PRESTADORES.get(d['cnes'], d['prestador']),
        COD_MUNICIPIO: d['municipio'],
        ORIGEM: d['origem'],
        'AIH': d['aih'],
        'Cod. Procedimento': d['codigo'],
        'Procedimento': d['descricao'],
        'Valor': converter_valor(d['valor']),
    }


def interpretar_linhas(linhas):
    texto = normalizar('\n'.join(linhas))
    if 'DEMONSTRATIVO DE PROCEDENCIA ANALITICO' not in texto:
        raise ValueError('Selecione o DEMONSTRATIVO DE PROCEDÊNCIA ANALÍTICO, em lugar dos relatórios individuais por estabelecimento.')
    competencias = {c for l in linhas if (c := extrair_competencia(l))}
    if len(competencias) != 1:
        raise ValueError('O PDF deve conter uma única competência identificável.')
    competencia = competencias.pop()
    if not 1 <= int(competencia[:2]) <= 12:
        raise ValueError('Competência inválida.')
    registros = []
    quantidade_bloco = 0
    valor_bloco = Decimal('0')
    qtd_conferida = False
    municipios_conferidos = 0
    totais_executantes = []
    for linha in linhas:
        simples = normalizar(linha)
        dados = interpretar_linha_producao(linha)
        if dados:
            if qtd_conferida:
                raise ValueError('Subtotal de município ausente antes do próximo registro.')
            registros.append({'ANO': competencia[-4:], 'MÊS': competencia[:2],
                              'Comp': competencia, **dados})
            quantidade_bloco += 1
            valor_bloco += dados['Valor']
        elif re.match(r'^\d{6,7}\s+', linha) or re.search(r'\b\d{13}\b', linha):
            raise ValueError(f'Linha de produção não reconhecida; a exportação foi interrompida: {linha}')
        elif simples.startswith('QTD TOTAL DE AIH'):
            qtd = re.fullmatch(r'QTD TOTAL DE AIHS?\s+(\d+)', simples)
            if not qtd or int(qtd.group(1)) != quantidade_bloco or qtd_conferida:
                raise ValueError('Quantidade de AIHs diverge do subtotal do município de origem.')
            qtd_conferida = True
        elif simples.startswith('TOTAL DO MUNICIPIO DE ORIGEM'):
            total = re.fullmatch(rf'TOTAL DO MUNICIPIO DE ORIGEM\s+({VALOR})', simples)
            if not total or not qtd_conferida or converter_valor(total.group(1)) != valor_bloco:
                raise ValueError('Valor extraído diverge do subtotal do município de origem.')
            municipios_conferidos += 1
            quantidade_bloco, valor_bloco, qtd_conferida = 0, Decimal('0'), False
        elif simples.startswith('TOTAL DO MUNICIPIO EXECUTANTE'):
            total = re.fullmatch(rf'TOTAL DO MUNICIPIO EXECUTANTE\s+(\d+)\s+((?:{VALOR}\s+){{4}}{VALOR})', simples)
            if not total:
                raise ValueError('Total do município executante não reconhecido.')
            numeros = [converter_valor(v) for v in total.group(2).split()]
            if sum(numeros[:-1]) != numeros[-1]:
                raise ValueError('Total do município executante inconsistente.')
            totais_executantes.append((int(total.group(1)), numeros[-1]))
    if not registros or not municipios_conferidos or quantidade_bloco or qtd_conferida:
        raise ValueError('PDF sem registros válidos ou com subtotais incompletos.')
    if (not totais_executantes
            or sum(q for q, v in totais_executantes) != len(registros)
            or sum(v for q, v in totais_executantes) != sum(r['Valor'] for r in registros)):
        raise ValueError('A extração não confere com os totais dos municípios executantes.')
    return pd.DataFrame(registros, columns=COLUNAS)


def processar_pdf(caminho_pdf):
    try:
        return interpretar_linhas(extrair_texto_pdf(caminho_pdf))
    except ValueError as erro:
        raise ValueError(f'{Path(caminho_pdf).name}: {erro}') from erro


def consolidar(arquivos_pdf):
    if not arquivos_pdf:
        raise ValueError('Nenhum PDF selecionado.')
    resultado = pd.concat([processar_pdf(p) for p in arquivos_pdf], ignore_index=True)
    if resultado.duplicated(subset=['Comp', 'CNES', 'AIH']).any():
        raise ValueError('Existem AIHs repetidas para a mesma competência e CNES. Verifique os PDFs selecionados.')
    return resultado.sort_values(['ANO', 'MÊS', 'CNES', 'AIH']).reset_index(drop=True)


def salvar_excel(resultado, arquivo_saida):
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter
    destino = Path(arquivo_saida)
    if destino.suffix.lower() != '.xlsx':
        raise ValueError('O arquivo de saída deve ter extensão .xlsx.')
    saida = resultado.copy()
    saida['Valor'] = saida['Valor'].map(float)
    with pd.ExcelWriter(destino, engine='openpyxl') as writer:
        saida.to_excel(writer, sheet_name='AIHs', index=False)
        planilha = writer.book['AIHs']
        for coluna in ['ANO', 'MÊS', 'Comp', 'CNES', COD_MUNICIPIO, ORIGEM, 'AIH', 'Cod. Procedimento']:
            for celula in planilha[get_column_letter(COLUNAS.index(coluna) + 1)][1:]:
                celula.number_format = '@'
        for celula in planilha[get_column_letter(COLUNAS.index('Valor') + 1)][1:]:
            celula.number_format = '#,##0.00'
        for n, largura in enumerate([8, 8, 12, 12, 32, 18, 42, 20, 22, 52, 18], 1):
            planilha.column_dimensions[get_column_letter(n)].width = largura
        for linha in planilha.iter_rows(min_row=2):
            for celula in (linha[COLUNAS.index(nome)] for nome in ['Prestador', ORIGEM, 'Procedimento']):
                celula.alignment = Alignment(vertical='center', wrap_text=True)
            planilha.row_dimensions[linha[0].row].height = 30
        planilha.freeze_panes = 'A2'
        planilha.auto_filter.ref = planilha.dimensions


def interface():
    import tkinter as tk
    from tkinter import filedialog, messagebox
    root = tk.Tk()
    root.withdraw()
    try:
        messagebox.showinfo('Extração PMAE por procedência',
            'Selecione o DEMONSTRATIVO DE PROCEDÊNCIA ANALÍTICO do SIHD.\n\n'
            'Gere o relatório com o filtro PMAE desejado. Todos os registros '
            'do PDF serão extraídos, com o código do município e o município/UF '
            'de origem em colunas separadas antes da AIH.\n\n'
            'É possível selecionar vários PDFs de competências distintas.')
        arquivos = filedialog.askopenfilenames(title='Selecione os demonstrativos de procedência analíticos', filetypes=[('Arquivos PDF', '*.pdf')])
        if not arquivos:
            return
        resultado = consolidar(arquivos)
        destino = filedialog.asksaveasfilename(title='Salvar resultado', initialfile='SIHD_AIHs_Procedimentos.xlsx', defaultextension='.xlsx', filetypes=[('Arquivo Excel', '*.xlsx')])
        if not destino:
            return
        salvar_excel(resultado, destino)
        messagebox.showinfo('Processamento concluído', f'Arquivo Excel gerado:\n{destino}\n\nRegistros: {len(resultado)}\nMunicípios de origem: {resultado[ORIGEM].nunique()}\nQuantidades e valores conferidos com o PDF.')
    except Exception as erro:
        messagebox.showerror('Não foi possível concluir a extração', str(erro))
    finally:
        root.destroy()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('pdfs', nargs='*')
    parser.add_argument('--saida')
    args = parser.parse_args()
    if not args.pdfs:
        if args.saida:
            parser.error('Informe o PDF de entrada.')
        interface()
        return
    if not args.saida:
        parser.error('Informe --saida arquivo.xlsx.')
    if Path(args.saida).exists():
        parser.error('O destino já existe. Escolha outro nome para preservar o arquivo.')
    resultado = consolidar(args.pdfs)
    salvar_excel(resultado, args.saida)
    print(f'Gerado: {args.saida} | {len(resultado)} registros | Total: {resultado["Valor"].sum():.2f}')


if __name__ == '__main__':
    main()
