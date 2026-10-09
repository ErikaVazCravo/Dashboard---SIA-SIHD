"""Converte o resumo SIHD2 por CNES para o formato SIHD_Prestadores.

Sem argumentos: interface de seleção de PDFs e destino.
Terminal: python "Resumo da Produção Hospitalar.py" relatorio.pdf --saida resultado.xlsx
Dependências: python -m pip install pdfplumber openpyxl
"""

import argparse
from copy import copy
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path
import re
import sys
import unicodedata


CABECALHOS = ["ANO", "MÊS", "COMP", "CNES", "Prestador", "Qtde. Apres.",
              "Qtde. Aprov", "Compl. Federal", "TOTAL"]
PRESTADORES = {
    "2079763": "CAIRBAR SCHUTEL",
    "6943284": "MATERNIDADE GOTA DE LEITE",
    "2082527": "SANTA CASA",
    "9267263": "UNIDADE DE RETAGUARDA",
}
INTEIRO = r"(?:\d{1,3}(?:\.\d{3})+|\d+)"
DINHEIRO = r"(?:\d{1,3}(?:\.\d{3})+|\d+),\d{2}"
INICIO = re.compile(rf"^(\d{{7}})\s*-\s*(.+?)\s+({INTEIRO})\s+({INTEIRO})\s+({DINHEIRO})$")


def normalizar(texto):
    return "".join(c for c in unicodedata.normalize("NFKD", texto)
                   if not unicodedata.combining(c)).lower()


def moeda(texto):
    return Decimal(texto.replace(".", "").replace(",", "."))


def copiar_estilo(origem, destino):
    # Índices de _style pertencem ao workbook de origem; copie os objetos.
    for atributo in ('font', 'fill', 'border', 'alignment', 'protection'):
        setattr(destino, atributo, copy(getattr(origem, atributo)))
    destino.number_format = origem.number_format


def valores(linha, quantidade):
    numeros = re.findall(DINHEIRO, linha)
    if len(numeros) != quantidade:
        raise ValueError(f"Linha com valores incompletos ou formato inesperado: {linha}")
    numeros = [moeda(n) for n in numeros]
    if sum(numeros[:-1]) != numeros[-1]:
        raise ValueError(f"Soma dos valores diverge do total: {linha}")
    return numeros[-1]


@dataclass(frozen=True)
class Registro:
    ano: int
    mes: int
    cnes: str
    nome: str
    apresentadas: int
    aprovadas: int
    federal: Decimal
    total: Decimal
    origem: str

    def linha(self):
        return [self.ano, self.mes, f"{self.mes:02d}/{self.ano}", self.cnes,
                PRESTADORES.get(self.cnes, self.nome), self.apresentadas,
                self.aprovadas, self.federal, self.total]


def interpretar(texto, origem="PDF"):
    """Falha em relatórios incompletos; não transforma campos ausentes em zero."""
    simples = normalizar(texto)
    if "resumo dos valores aprovados por cnes" not in simples:
        raise ValueError("O PDF não é um resumo dos valores aprovados por CNES do SIHD2.")
    competencias = set(re.findall(r"competencia\s*:\s*(\d{2})/(\d{4})", simples))
    if len(competencias) != 1:
        raise ValueError("É necessário haver uma única competência identificável por PDF.")
    mes, ano = map(int, next(iter(competencias)))
    if not 1 <= mes <= 12:
        raise ValueError("Mês da competência inválido.")

    registros, atual, campos, vistos = [], None, {}, set()

    def concluir():
        if atual is None:
            return
        cnes, nome, apresentadas, aprovadas, _ = atual.groups()
        if set(campos) != {"base", "federal", "local", "total"}:
            raise ValueError(f"CNES {cnes}: bloco incompleto ({', '.join(campos)}).")
        if campos["base"] + campos["federal"] + campos["local"] != campos["total"]:
            raise ValueError(f"CNES {cnes}: total diverge dos valores e complementos.")
        if cnes in vistos:
            raise ValueError(f"CNES {cnes} repetido no PDF. Verifique o relatório.")
        ap, av = int(apresentadas.replace('.', '')), int(aprovadas.replace('.', ''))
        if av > ap:
            raise ValueError(f"CNES {cnes}: quantidade aprovada maior que apresentada.")
        vistos.add(cnes)
        registros.append(Registro(ano, mes, cnes, nome, ap, av,
                                  campos['federal'], campos['total'], origem))

    for linha in texto.splitlines():
        linha = " ".join(linha.split())
        chave = normalizar(linha)
        inicio = INICIO.fullmatch(linha)
        if inicio:
            concluir()
            atual, campos = inicio, {}
            continue
        if re.match(r"^\d{7}\s*-", linha):
            raise ValueError(f"Cabeçalho de prestador não reconhecido: {linha}")
        if chave.startswith(("total do municipio", "total do gestor")):
            concluir()
            atual, campos = None, {}
        if atual is None:
            continue
        for rotulo, campo, quantidade in [("valor sem complemento", "base", 5),
                                         ("complemento federal", "federal", 5),
                                         ("complemento local", "local", 5),
                                         ("total ", "total", 3)]:
            if chave.startswith(rotulo):
                if campo in campos:
                    raise ValueError(f"Linha repetida para o CNES {atual.group(1)}: {rotulo}")
                campos[campo] = valores(linha, quantidade)
                break
    concluir()
    if not registros:
        raise ValueError("Nenhum prestador encontrado. PDFs digitalizados precisam de OCR.")

    # O total do gestor inclui todos os prestadores; nunca é exportado como prestador.
    if "total do gestor" not in simples:
        raise ValueError("Total do gestor ausente: não é possível conferir a extração completa.")
    rodape = simples.rsplit("total do gestor", 1)[1]
    totais = re.findall(rf"(?:^|\n)\s*(?:total\s+)?({INTEIRO})\s+({INTEIRO})\s+{DINHEIRO}\s+({DINHEIRO})\s+({DINHEIRO})\s+({DINHEIRO})\s*(?:\n|$)", rodape)
    if len(totais) != 1:
        raise ValueError("Não foi possível identificar o total geral do gestor.")
    ap, av, hospitalar, profissional, total = totais[0]
    if moeda(hospitalar) + moeda(profissional) != moeda(total):
        raise ValueError("Total do gestor inconsistente.")
    if (sum(r.apresentadas for r in registros) != int(ap.replace('.', ''))
            or sum(r.aprovadas for r in registros) != int(av.replace('.', ''))
            or sum(r.total for r in registros) != moeda(total)):
        raise ValueError("A soma dos prestadores não confere com o total do gestor.")
    return registros


def extrair_pdf(caminho):
    import pdfplumber
    with pdfplumber.open(caminho) as pdf:
        paginas = [p.extract_text(x_tolerance=2) or "" for p in pdf.pages]
    if not paginas or any(not p.strip() for p in paginas):
        raise ValueError(f"{Path(caminho).name}: PDF vazio ou página sem texto; necessita OCR.")
    try:
        return interpretar("\n".join(paginas), str(Path(caminho).resolve()))
    except ValueError as erro:
        raise ValueError(f"{Path(caminho).name}: {erro}") from erro


def converter(pdfs, saida, modelo=None):
    from openpyxl import Workbook, load_workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    saida = Path(saida)
    if saida.suffix.lower() != '.xlsx':
        raise ValueError("O destino deve ter extensão .xlsx.")
    if saida.exists():
        raise FileExistsError(f"Já existe um arquivo no destino. Escolha outro nome: {saida}")
    registros, chaves = [], set()
    for pdf in pdfs:
        for r in extrair_pdf(pdf):
            chave = (r.ano, r.mes, r.cnes)
            if chave in chaves:
                raise ValueError(f"Competência/CNES duplicados entre PDFs: {r.mes:02d}/{r.ano}, {r.cnes}.")
            chaves.add(chave)
            registros.append(r)
    if not registros:
        raise ValueError("Selecione pelo menos um PDF.")
    ordem = {cnes: i for i, cnes in enumerate(PRESTADORES)}
    registros.sort(key=lambda r: (r.ano, r.mes, ordem.get(r.cnes, len(ordem)), r.nome))
    wb = Workbook()
    ws = wb.active
    ws.title = "SIHD_Prestadores"
    ws.append(CABECALHOS)
    referencia = None
    if modelo:
        referencia = load_workbook(modelo)
        ref = referencia['SIHD_Prestadores']
        if [ref.cell(1, c).value for c in range(1, 10)] != CABECALHOS:
            raise ValueError("O modelo não possui os nove cabeçalhos esperados.")
    larguras = [8.71, 7.14, 13.14, 11.43, 42.29, 17, 17, 17, 19]
    for c, largura in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(c)].width = largura
        celula = ws.cell(1, c)
        celula.font = Font(name="Calibri", size=11, bold=True)
        celula.alignment = Alignment(horizontal="center", vertical="center")
        if referencia:
            copiar_estilo(ref.cell(1, c), celula)
    ws.row_dimensions[1].height = 24
    for n, r in enumerate(registros, 2):
        ws.append(r.linha())
        ws.row_dimensions[n].height = 20
        for c in range(1, 10):
            celula = ws.cell(n, c)
            celula.font = Font(name="Calibri", size=11)
            celula.alignment = Alignment(horizontal="left" if c == 5 else "center" if c < 5 else "right", vertical="center")
            if r.mes % 2 == 0:
                celula.fill = PatternFill("solid", fgColor="CCCCCC")
            if referencia:
                copiar_estilo(ref.cell(2 if r.mes % 2 else 6, c), celula)
            celula.number_format = ('"R$" #,##0.00' if c >= 8 else '@' if c in (3, 4, 5) else '0')
            if c == 5:
                celula.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                if len(str(celula.value)) > 38:
                    ws.row_dimensions[n].height = 34
        ws.cell(n, 4).comment = Comment(f"Fonte: {r.origem}\nCompetência: {r.mes:02d}/{r.ano}\nCNES: {r.cnes}\nNome no PDF: {r.nome}", "Automação SIHD")
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:I{ws.max_row}'
    ws.print_title_rows = '1:1'
    ws.print_area = f'A1:I{ws.max_row}'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    saida.parent.mkdir(parents=True, exist_ok=True)
    # Criação exclusiva: impede substituir o modelo ou uma produção anterior.
    with saida.open('xb') as arquivo:
        wb.save(arquivo)
    if referencia:
        referencia.close()
    return registros


def interface():
    import tkinter as tk
    from tkinter import filedialog, messagebox
    root = tk.Tk()
    root.withdraw()
    try:
        pdfs = filedialog.askopenfilenames(title="Resumo da Produção Hospitalar - selecione os PDFs SIHD2", filetypes=[("Relatórios PDF", "*.pdf")])
        if not pdfs:
            return
        saida = filedialog.asksaveasfilename(title="Salvar Resumo da Produção Hospitalar", initialfile="Resumo da Produção Hospitalar.xlsx", defaultextension=".xlsx", filetypes=[("Planilha Excel", "*.xlsx")])
        if not saida:
            return
        registros = converter(pdfs, saida)
        messagebox.showinfo("Resumo da Produção Hospitalar", f"Planilha gerada e totais conferidos.\n\n{len(registros)} registros.\n\n{saida}")
    except Exception as erro:
        messagebox.showerror("Não foi possível converter", str(erro))
    finally:
        root.destroy()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('pdfs', nargs='*', type=Path)
    parser.add_argument('--saida', type=Path)
    parser.add_argument('--modelo', type=Path, help='Modelo opcional para copiar os estilos; não importa o histórico.')
    args = parser.parse_args()
    if not args.pdfs:
        if args.saida or args.modelo:
            parser.error('Informe os PDFs para usar --saida ou --modelo.')
        interface()
        return
    if not args.saida:
        parser.error('Informe --saida resultado.xlsx.')
    try:
        registros = converter(args.pdfs, args.saida, args.modelo)
    except Exception as erro:
        print(f'Erro: {erro}', file=sys.stderr)
        sys.exit(1)
    print(f'Gerado: {args.saida} | {len(registros)} registros | Totais conferidos.')


if __name__ == '__main__':
    main()
