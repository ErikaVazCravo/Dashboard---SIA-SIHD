import os
import re
import sys
import math
import pandas as pd
import tkinter as tk
from tkinter import filedialog, messagebox
from openpyxl.styles import PatternFill


# ============================================================
# PARTE 1/4
# ESTRUTURA GERAL, INTERFACE E FUNÇÕES AUXILIARES
# ============================================================


# ============================================================
# INTERFACE INICIAL
# ============================================================

root = tk.Tk()
root.withdraw()

messagebox.showinfo(
    "OCI - Produção Prestadores",
    "AUTOMAÇÃO OCI - PRODUÇÃO PRESTADORES\n\n"
    "Este sistema irá executar o processamento completo das OCIs.\n\n"
    "Fluxo da automação:\n\n"
    "1. O programa utilizará automaticamente a pasta SIGTAP localizada no mesmo diretório do arquivo .py.\n"
    "   Essa pasta deverá conter:\n"
    "   - tb_procedimento.txt\n"
    "   - rl_procedimento_compativel.txt\n"
    "2. Gerar a base de valores dos procedimentos OCI.\n\n"
    "3. Selecionar o arquivo SIA/APAC.\n\n"
    "4. Utilizar o arquivo Nomes_estabelecimentos.txt, se disponível.\n\n"
    "5. Gerar um Excel final com as abas:\n"
    "   - DADOS_OCI\n"
    "   - Vl_SIGTAP\n"
    "   - Vl_APAC_SIGTAP\n"
    "   - Total_Terceiros\n\n"
    "Observação:\n"
    "Processamento corresponde à competência em que a produção foi processada no SIA.\n"
    "Comp. corresponde à competência de produção informada no relatório."
)


# ============================================================
# CAMINHOS E SELEÇÕES
# ============================================================

def pasta_do_programa():
    """
    Retorna a pasta onde está o arquivo .py ou o executável.
    """
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)

    return os.path.dirname(os.path.abspath(__file__))


def selecionar_pasta(titulo):
    """
    Abre janela para seleção de pasta.
    """
    return filedialog.askdirectory(title=titulo)


def selecionar_arquivo(titulo):
    """
    Abre janela para seleção de arquivo.
    """
    return filedialog.askopenfilename(
        title=titulo,
        filetypes=[
            ("Todos os arquivos", "*.*")
        ]
    )


def salvar_arquivo_excel(titulo, nome_padrao="Resultado_OCI.xlsx"):
    """
    Abre janela para definir o local de salvamento do Excel final.
    """
    return filedialog.asksaveasfilename(
        title=titulo,
        initialfile=nome_padrao,
        defaultextension=".xlsx",
        filetypes=[
            ("Arquivo Excel", "*.xlsx"),
            ("Todos os arquivos", "*.*")
        ]
    )


# ============================================================
# LIMPEZA E NORMALIZAÇÃO
# ============================================================

def limpar_linha(linha):
    """
    Limpa caracteres especiais comuns nos relatórios do SIA,
    mantendo os espaços necessários para leitura por posição fixa.
    """
    if linha is None:
        return ""

    return (
        str(linha)
        .replace("\x0c", "")
        .replace("\xa0", " ")
        .rstrip("\r\n")
    )


def limpar_texto_excel(valor):
    """
    Remove caracteres ilegais para gravação em Excel.
    """
    if isinstance(valor, str):
        return re.sub(r"[\x00-\x1F]", "", valor)

    return valor


def limpar_valor(valor):
    """
    Prepara valores monetários no padrão brasileiro para conversão.
    Exemplo:
    1.234,56 -> 1234.56
    """
    if valor is None:
        return ""

    return (
        str(valor)
        .replace(".", "")
        .replace(",", ".")
        .strip()
    )


def normalizar_cnes(valor):
    """
    Normaliza CNES para 7 dígitos.
    """
    if valor is None or pd.isna(valor):
        return ""

    digitos = re.sub(r"\D", "", str(valor))

    if not digitos:
        return ""

    if len(digitos) < 7:
        return digitos.zfill(7)

    return digitos[:7]


def normalizar_proc(valor):
    """
    Normaliza código de procedimento para 10 dígitos.

    Trata:
    - valores numéricos do Excel;
    - códigos com pontuação;
    - códigos sem zero à esquerda.
    """
    if valor is None or (isinstance(valor, float) and pd.isna(valor)):
        return ""

    if isinstance(valor, (int, float)):
        try:
            return str(int(float(valor))).zfill(10)
        except:
            return ""

    texto = str(valor).strip()

    if not texto:
        return ""

    if re.fullmatch(r"\d+\.0+", texto):
        texto = texto.split(".")[0]

    digitos = re.sub(r"\D", "", texto)

    if not digitos:
        return ""

    if len(digitos) < 10:
        return digitos.zfill(10)

    if len(digitos) > 10:
        return digitos[:10]

    return digitos


# ============================================================
# CONVERSÕES NUMÉRICAS
# ============================================================

def parse_numero_inteiro(valor):
    """
    Converte texto para inteiro.
    """
    texto = re.sub(r"\D", "", str(valor or ""))

    if not texto:
        return 0

    try:
        return int(texto)
    except:
        return 0


def parse_numero_real(valor):
    """
    Converte texto monetário para float.
    """
    texto = limpar_valor(valor)

    if not texto:
        return 0.0

    try:
        return float(texto)
    except:
        return 0.0


def parse_valor_sigtap(valor):
    """
    Converte valor da tabela SIGTAP.
    No SIGTAP, valores monetários geralmente vêm sem vírgula,
    com centavos nas duas últimas casas.
    Exemplo:
    000000014000 -> 140.00
    """
    valor_num = pd.to_numeric(valor, errors="coerce")

    if pd.isna(valor_num):
        return 0.0

    return round(float(valor_num) / 100, 2)


# ============================================================
# COMPETÊNCIAS E DATAS
# ============================================================

def formatar_competencia_sigtap(valor):
    """
    Converte competência AAAAMM em MM/AAAA.
    Exemplo:
    202605 -> 05/2026
    """
    valor = str(valor or "").strip()

    if len(valor) == 6 and valor.isdigit():
        return f"{valor[4:6]}/{valor[0:4]}"

    return valor


def formatar_mm_aaaa(valor):
    """
    Garante formato MM/AAAA para valores que já estejam nesse padrão
    ou que tenham sido interpretados como data pelo pandas/Excel.
    """
    if valor is None or pd.isna(valor):
        return ""

    texto = str(valor).strip()

    if re.fullmatch(r"\d{2}/\d{4}", texto):
        return texto

    data = pd.to_datetime(texto, errors="coerce")

    if pd.isna(data):
        return texto

    return data.strftime("%m/%Y")


def formatar_coluna_mm_aaaa(df, coluna):
    """
    Aplica formato MM/AAAA em uma coluna, se ela existir.
    """
    if df.empty or coluna not in df.columns:
        return df

    df = df.copy()
    df[coluna] = df[coluna].apply(formatar_mm_aaaa)

    return df


# ============================================================
# EXTRAÇÃO POR POSIÇÃO FIXA
# ============================================================

def extrair_faixa(linha, inicio, fim):
    """
    Extrai trecho por posição fixa 1-indexada e inclusiva.

    Exemplo:
    inicio=2, fim=3 extrai caracteres das posições 2 e 3.
    """
    if not linha:
        return ""

    inicio_idx = max(inicio - 1, 0)
    fim_idx = min(fim, len(linha))

    if inicio_idx >= len(linha):
        return ""

    return linha[inicio_idx:fim_idx].strip()


# ============================================================
# REGRAS DE OCI
# ============================================================

def eh_apac_oci_por_numero(numero_apac):
    """
    Regra OCI:
    - APAC inicia com 35;
    - 5º dígito numérico da APAC é 6 ou 7.
    """
    if not numero_apac:
        return False

    digitos = re.sub(r"\D", "", str(numero_apac))

    if len(digitos) < 5:
        return False

    return digitos.startswith("35") and digitos[4] in ("6", "7")


# ============================================================
# MAPAMENTOS
# ============================================================

def mapear_tipo_registro(codigo):
    """
    Mapeia tipo de registro da tabela SIGTAP.
    """
    return {
        "06": "APAC (Proc. Principal)",
        "07": "APAC (Proc. Secundário)"
    }.get(str(codigo).strip(), "")


def mapear_compatibilidade(codigo):
    """
    Mapeia tipo de compatibilidade do SIGTAP.
    """
    return {
        "1": "Compatível",
        "2": "Incompatível / Excludente",
        "3": "Concomitante",
        "5": "Obrigatório"
    }.get(str(codigo).strip(), "")


# ============================================================
# ARQUIVO DE NOMES DE ESTABELECIMENTOS
# ============================================================

def carregar_mapa_estabelecimentos():
    """
    Lê o arquivo Nomes_estabelecimentos.txt.

    Primeiro procura no mesmo local do programa.
    Se não encontrar, permite seleção manual.

    Formato esperado:
    CNES com 7 dígitos + TAB + Nome do estabelecimento
    """
    caminho_txt = os.path.join(
        pasta_do_programa(),
        "Nomes_estabelecimentos.txt"
    )

    if not os.path.exists(caminho_txt):
        messagebox.showinfo(
            "Nomes_estabelecimentos.txt",
            "O arquivo Nomes_estabelecimentos.txt não foi encontrado "
            "no mesmo local do programa.\n\n"
            "Selecione o arquivo manualmente."
        )

        caminho_txt = filedialog.askopenfilename(
            title="Selecione o arquivo Nomes_estabelecimentos.txt",
            filetypes=[
                ("Arquivo TXT", "*.txt"),
                ("Todos os arquivos", "*.*")
            ]
        )

    mapa = {}

    if not caminho_txt or not os.path.exists(caminho_txt):
        print("⚠️ Arquivo Nomes_estabelecimentos.txt não selecionado.")
        return mapa

    with open(caminho_txt, "r", encoding="utf-8-sig", errors="replace") as arquivo:
        for numero_linha, linha in enumerate(arquivo, start=1):
            linha = linha.rstrip("\n")

            if not linha.strip():
                continue

            partes = linha.split("\t", maxsplit=1)

            if len(partes) != 2:
                print(
                    f"⚠️ Linha ignorada em Nomes_estabelecimentos.txt "
                    f"(sem TAB): linha {numero_linha} -> {repr(linha)}"
                )
                continue

            cnes = normalizar_cnes(partes[0])
            nome = partes[1].strip()

            if len(cnes) == 7 and nome:
                mapa[cnes] = nome
            else:
                print(
                    f"⚠️ Linha ignorada em Nomes_estabelecimentos.txt "
                    f"(CNES inválido ou nome vazio): linha {numero_linha} -> {repr(linha)}"
                )

    print(f"✅ Nomes de estabelecimentos carregados: {len(mapa)}")
    print(f"📄 Arquivo de nomes utilizado: {caminho_txt}")

    return mapa


def preencher_nomes_estabelecimentos(df, mapa_cnes):
    """
    Substitui nomes dos estabelecimentos conforme Nomes_estabelecimentos.txt.
    """
    if df.empty or not mapa_cnes:
        return df

    df = df.copy()

    if "CNES Estab. Exec." in df.columns:
        cnes_exec = df["CNES Estab. Exec."].apply(normalizar_cnes)
        df["Nome Estab. Exec."] = cnes_exec.map(mapa_cnes).fillna(
            df.get("Nome Estab. Exec.", "")
        )

    if "Cnes Terc." in df.columns:
        cnes_terc = df["Cnes Terc."].apply(normalizar_cnes)
        nomes_terc = cnes_terc.map(mapa_cnes).fillna("")

        if "Nome Terc." in df.columns:
            df["Nome Terc."] = df["Nome Terc."].astype(str).str.strip()
            df["Nome Terc."] = df["Nome Terc."].where(
                df["Nome Terc."] != "",
                nomes_terc
            )
        else:
            df["Nome Terc."] = nomes_terc

    return df
# ============================================================
# PARTE 2/4
# LEITURA SIGTAP E GERAÇÃO DA BASE DE VALORES OCI
# ============================================================


# ============================================================
# LEITURA DAS LINHAS DO SIGTAP
# ============================================================

def ler_linha_rl_procedimento_compativel(linha):
    """
    Lê uma linha do arquivo rl_procedimento_compativel.txt.

    Layout:
    CO_PROCEDIMENTO_PRINCIPAL   1-10
    CO_REGISTRO_PRINCIPAL       11-12
    CO_PROCEDIMENTO_COMPATIVEL  13-22
    CO_REGISTRO_COMPATIVEL      23-24
    TP_COMPATIBILIDADE          25
    QT_PERMITIDA                26-29
    DT_COMPETENCIA              30-35
    """
    return {
        "CO_PROCEDIMENTO_PRINCIPAL": linha[0:10].strip(),
        "CO_REGISTRO_PRINCIPAL": linha[10:12].strip(),
        "CO_PROCEDIMENTO_COMPATIVEL": linha[12:22].strip(),
        "CO_REGISTRO_COMPATIVEL": linha[22:24].strip(),
        "TP_COMPATIBILIDADE": linha[24:25].strip(),
        "QT_PERMITIDA": linha[25:29].strip(),
        "DT_COMPETENCIA": linha[29:35].strip(),
    }


def ler_linha_tb_procedimento(linha):
    """
    Lê uma linha do arquivo tb_procedimento.txt.

    Layout:
    CO_PROCEDIMENTO   1-10
    NO_PROCEDIMENTO   11-260
    VL_SA             295-306
    DT_COMPETENCIA    331-336
    """
    return {
        "CO_PROCEDIMENTO": linha[0:10].strip(),
        "NO_PROCEDIMENTO": linha[10:260].strip(),
        "VL_SA": linha[294:306].strip(),
        "DT_COMPETENCIA": linha[330:336].strip(),
    }


# ============================================================
# VALIDAÇÃO DA PASTA SIGTAP
# ============================================================

def validar_pasta_sigtap(pasta_sigtap):
    """
    Verifica se a pasta SIGTAP contém os arquivos necessários.
    """
    caminho_rl = os.path.join(
        pasta_sigtap,
        "rl_procedimento_compativel.txt"
    )

    caminho_tb = os.path.join(
        pasta_sigtap,
        "tb_procedimento.txt"
    )

    erros = []

    if not os.path.exists(caminho_rl):
        erros.append(f"rl_procedimento_compativel.txt não encontrado:\n{caminho_rl}")

    if not os.path.exists(caminho_tb):
        erros.append(f"tb_procedimento.txt não encontrado:\n{caminho_tb}")

    if erros:
        messagebox.showerror(
            "Arquivos SIGTAP não encontrados",
            "\n\n".join(erros)
        )
        raise SystemExit("Arquivos SIGTAP obrigatórios não encontrados.")

    return caminho_rl, caminho_tb


# ============================================================
# CARREGAMENTO DO SIGTAP
# ============================================================

def carregar_rl_procedimento_compativel(caminho_rl):
    """
    Carrega o arquivo rl_procedimento_compativel.txt.
    Mantém apenas procedimentos principais iniciados por 09.
    """
    dados = []

    with open(caminho_rl, "r", encoding="latin-1", errors="ignore") as arquivo:
        for linha in arquivo:
            if not linha.strip():
                continue

            dados.append(ler_linha_rl_procedimento_compativel(linha))

    df = pd.DataFrame(dados)

    if df.empty:
        return df

    df["CO_PROCEDIMENTO_PRINCIPAL"] = df["CO_PROCEDIMENTO_PRINCIPAL"].apply(normalizar_proc)
    df["CO_PROCEDIMENTO_COMPATIVEL"] = df["CO_PROCEDIMENTO_COMPATIVEL"].apply(normalizar_proc)

    df = df[
        df["CO_PROCEDIMENTO_PRINCIPAL"].astype(str).str.startswith("09")
    ].copy()

    return df


def carregar_tb_procedimento(caminho_tb):
    """
    Carrega o arquivo tb_procedimento.txt.
    """
    dados = []

    with open(caminho_tb, "r", encoding="latin-1", errors="ignore") as arquivo:
        for linha in arquivo:
            if not linha.strip():
                continue

            dados.append(ler_linha_tb_procedimento(linha))

    df = pd.DataFrame(dados)

    if df.empty:
        return df

    df["CO_PROCEDIMENTO"] = df["CO_PROCEDIMENTO"].apply(normalizar_proc)
    df["NO_PROCEDIMENTO"] = df["NO_PROCEDIMENTO"].astype(str).str.strip()
    df["VL_SA"] = df["VL_SA"].apply(parse_valor_sigtap)
    df["DT_COMPETENCIA"] = df["DT_COMPETENCIA"].astype(str).str.strip()

    return df


# ============================================================
# GERAÇÃO DA BASE OCI PROCEDIMENTOS E VALORES
# ============================================================

def gerar_base_oci_procedimentos_e_valores(pasta_sigtap):
    """
    Gera a base de procedimentos OCI a partir dos arquivos SIGTAP.

    Retorna:
    - df_resultado
    - caminho do Excel gerado

    Também salva:
    OCI_Procedimentos_e_Valores.xlsx
    dentro da pasta SIGTAP.
    """
    caminho_rl, caminho_tb = validar_pasta_sigtap(pasta_sigtap)

    print("\n========================================")
    print(" GERANDO BASE OCI DO SIGTAP")
    print("========================================")

    print(f"📄 Lendo: {caminho_rl}")
    df_rl = carregar_rl_procedimento_compativel(caminho_rl)

    print(f"📄 Lendo: {caminho_tb}")
    df_tb = carregar_tb_procedimento(caminho_tb)

    if df_rl.empty:
        messagebox.showwarning(
            "SIGTAP sem OCI",
            "Nenhum procedimento principal iniciado por 09 foi localizado "
            "em rl_procedimento_compativel.txt."
        )

    if df_tb.empty:
        messagebox.showerror(
            "Erro",
            "O arquivo tb_procedimento.txt não possui dados válidos."
        )
        raise SystemExit("tb_procedimento.txt sem dados válidos.")

    # --------------------------------------------------------
    # JOIN COM PROCEDIMENTO PRINCIPAL
    # --------------------------------------------------------

    df = df_rl.merge(
        df_tb,
        left_on="CO_PROCEDIMENTO_PRINCIPAL",
        right_on="CO_PROCEDIMENTO",
        how="left",
        suffixes=("", "_TB_PRINC")
    )

    df = df.rename(columns={
        "NO_PROCEDIMENTO": "Nome_Proc_Principal",
        "VL_SA": "Vl_Proc_Principal",
        "DT_COMPETENCIA_TB_PRINC": "DT_COMPETENCIA_TB_PRINC"
    })

    if "CO_PROCEDIMENTO" in df.columns:
        df = df.drop(columns=["CO_PROCEDIMENTO"])

    # --------------------------------------------------------
    # JOIN COM PROCEDIMENTO SECUNDÁRIO / COMPATÍVEL
    # --------------------------------------------------------

    df = df.merge(
        df_tb,
        left_on="CO_PROCEDIMENTO_COMPATIVEL",
        right_on="CO_PROCEDIMENTO",
        how="left",
        suffixes=("", "_TB_SEC")
    )

    df = df.rename(columns={
        "NO_PROCEDIMENTO": "Nome_Proc_Secundario",
        "VL_SA": "Vl_Proc_Secundario",
    })

    if "CO_PROCEDIMENTO" in df.columns:
        df = df.drop(columns=["CO_PROCEDIMENTO"])

    # --------------------------------------------------------
    # RESULTADO FINAL
    # --------------------------------------------------------

    resultado = pd.DataFrame({
        "Comp.": df["DT_COMPETENCIA"].apply(formatar_competencia_sigtap),
        "Proc_Principal": df["CO_PROCEDIMENTO_PRINCIPAL"],
        "Nome_Proc_Principal": df["Nome_Proc_Principal"],
        "Vl_Proc_Principal": df["Vl_Proc_Principal"],
        "Cod_Instr_Registro_PRINC": df["CO_REGISTRO_PRINCIPAL"],
        "Tp_Instr_Registro_PRINC": df["CO_REGISTRO_PRINCIPAL"].apply(mapear_tipo_registro),
        "Proc_Secundario": df["CO_PROCEDIMENTO_COMPATIVEL"],
        "Nome_Proc_Secundario": df["Nome_Proc_Secundario"],
        "Vl_Proc_Secundario": df["Vl_Proc_Secundario"],
        "Cod_Instr_Registro_SEC": df["CO_REGISTRO_COMPATIVEL"],
        "Tp_Instr_Registro_SEC": df["CO_REGISTRO_COMPATIVEL"].apply(mapear_tipo_registro),
        "Cod_Compatibilidade": df["TP_COMPATIBILIDADE"],
        "Tp_Compatibilidade": df["TP_COMPATIBILIDADE"].apply(mapear_compatibilidade),
        "Qtd_Max": pd.to_numeric(df["QT_PERMITIDA"], errors="coerce"),
    })

    resultado = resultado.map(limpar_texto_excel)

    caminho_saida = os.path.join(
        pasta_sigtap,
        "OCI_Procedimentos_e_Valores.xlsx"
    )

    resultado.to_excel(caminho_saida, index=False)

    print("✅ Base OCI SIGTAP gerada com sucesso.")
    print(f"📄 Arquivo gerado: {caminho_saida}")
    print(f"📊 Total de linhas: {len(resultado)}")

    return resultado, caminho_saida


# ============================================================
# MAPAS A PARTIR DA BASE OCI
# ============================================================

def carregar_mapas_valores_sigtap(df_sigtap):
    """
    Cria mapas de valores:
    - procedimento principal -> valor principal
    - procedimento principal + procedimento secundário -> valor secundário
    """
    principal_map = {}
    secundario_map = {}

    if df_sigtap.empty:
        return principal_map, secundario_map

    for _, row in df_sigtap.iterrows():
        proc_principal = normalizar_proc(row.get("Proc_Principal", ""))
        proc_secundario = normalizar_proc(row.get("Proc_Secundario", ""))

        vl_principal = pd.to_numeric(
            row.get("Vl_Proc_Principal"),
            errors="coerce"
        )

        vl_secundario = pd.to_numeric(
            row.get("Vl_Proc_Secundario"),
            errors="coerce"
        )

        if proc_principal and proc_principal not in principal_map and pd.notna(vl_principal):
            principal_map[proc_principal] = float(vl_principal)

        if (
            proc_principal
            and proc_secundario
            and (proc_principal, proc_secundario) not in secundario_map
            and pd.notna(vl_secundario)
        ):
            secundario_map[(proc_principal, proc_secundario)] = float(vl_secundario)

    return principal_map, secundario_map


def carregar_mapas_nomes_procedimentos(df_sigtap):
    """
    Cria mapas de nomes de procedimentos:
    - procedimento principal -> nome principal
    - procedimento secundário -> nome secundário
    """
    mapa_principal = {}
    mapa_secundario = {}

    if df_sigtap.empty:
        return mapa_principal, mapa_secundario

    for _, row in df_sigtap.iterrows():
        proc_principal = normalizar_proc(row.get("Proc_Principal", ""))
        nome_principal = str(row.get("Nome_Proc_Principal", "")).strip()

        proc_secundario = normalizar_proc(row.get("Proc_Secundario", ""))
        nome_secundario = str(row.get("Nome_Proc_Secundario", "")).strip()

        if proc_principal:
            mapa_principal[proc_principal] = nome_principal

        if proc_secundario:
            mapa_secundario[proc_secundario] = nome_secundario

    return mapa_principal, mapa_secundario

# ============================================================
# PARTE 3/4
# LEITURA DO RELATÓRIO SIA/APAC
# ============================================================


# ============================================================
# EXTRAÇÕES DO CABEÇALHO DA APAC
# ============================================================

def extrair_cnes_nome_estab(linha):
    """
    Extrai CNES e nome do estabelecimento executor.

    Exemplo:
    Estab.Exec. : 2042487 - DAVITA
    """
    m = re.search(
        r"Estab\.Exec\.\s*:\s*(\d{7})\s*-\s*(.*?)(?:\s+Estab\.Solic\.|$)",
        linha
    )

    if m:
        cnes = normalizar_cnes(m.group(1))
        nome = m.group(2).strip()
        return cnes, nome

    return "", ""


def extrair_apac(linha):
    """
    Extrai número da APAC.

    Exemplo:
    APAC: 352623960493-2
    """
    m = re.search(r"APAC:\s*([0-9\-]{14})", linha)

    if m:
        return m.group(1).strip()

    return ""


def extrair_processamento(linha):
    """
    Extrai Processamento.

    Exemplo:
    Processamento:05/2026
    """
    m = re.search(r"Processamento:\s*(\d{2}/\d{4})", linha)

    if m:
        return m.group(1).strip()

    return ""


def extrair_producao(linha):
    """
    Extrai Produção, que será gravada como Comp.

    Exemplo:
    Producao:05/2026
    """
    m = re.search(r"Producao:\s*(\d{2}/\d{4})", linha)

    if m:
        return m.group(1).strip()

    return ""


def extrair_paciente(linha):
    """
    Extrai CPF, CNS e nome do paciente.

    Regras:
    - Após 'Paciente :' pode haver CPF de 11 dígitos;
    - Pode haver CNS de 15 dígitos;
    - Nome vem após o hífen.
    """
    cpf = ""
    cns = ""
    nome = ""

    m = re.search(r"Paciente\s*:\s*(.*?)\s*-\s*(.*)$", linha)

    if not m:
        return cpf, cns, nome

    bloco_documentos = m.group(1).strip()
    nome = m.group(2).strip()

    numeros = re.findall(r"\d+", bloco_documentos)

    if numeros:
        primeiro = numeros[0]

        if len(primeiro) == 11:
            cpf = primeiro

            if len(numeros) > 1 and len(numeros[1]) == 15:
                cns = numeros[1]

        elif len(primeiro) == 15:
            cns = primeiro

    return cpf, cns, nome


def extrair_data_nasc(linha):
    """
    Extrai data de nascimento.

    Exemplo:
    Dt.Nasc.: 01/09/1940
    """
    m = re.search(r"Dt\.Nasc\.\s*:\s*(\d{2}/\d{2}/\d{4})", linha)

    if m:
        return m.group(1)

    return ""


# ============================================================
# PARSER PRINCIPAL DO ARQUIVO APAC
# ============================================================

def parse_apac(caminho_arquivo):
    """
    Lê o arquivo SIA/APAC e extrai somente APACs OCI.

    Critério OCI:
    - APAC inicia com 35;
    - 5º dígito numérico da APAC é 6 ou 7.

    Retorna:
    - DataFrame com procedimentos;
    - total de APACs OCI identificadas;
    - total de APACs OCI com procedimentos extraídos.
    """
    with open(caminho_arquivo, "r", encoding="latin-1", errors="ignore") as arquivo:
        linhas = arquivo.readlines()

    dados = []
    chaves_registro = set()
    apacs_oci_identificadas = set()
    apacs_oci_com_procedimentos = set()

    i = 0

    while i < len(linhas):
        linha = limpar_linha(linhas[i])
        numero_apac_linha = extrair_apac(linha)

        if numero_apac_linha and eh_apac_oci_por_numero(numero_apac_linha):
            apacs_oci_identificadas.add(numero_apac_linha)

            cnes = ""
            nome_estab = ""
            numero_apac = numero_apac_linha
            processamento = ""
            comp = ""
            cpf = ""
            cns = ""
            nome_pac = ""
            data_nasc = ""

            # ------------------------------------------------
            # Busca inicial para cima
            # ------------------------------------------------

            if i - 3 >= 0:
                cnes, nome_estab = extrair_cnes_nome_estab(
                    limpar_linha(linhas[i - 3])
                )

            for j in range(i, max(i - 20, -1), -1):
                linha_up = limpar_linha(linhas[j])

                if (not cnes or not nome_estab) and "Estab.Exec." in linha_up:
                    cnes_tmp, nome_tmp = extrair_cnes_nome_estab(linha_up)
                    cnes = cnes or cnes_tmp
                    nome_estab = nome_estab or nome_tmp

                if not numero_apac and "APAC:" in linha_up:
                    numero_apac = extrair_apac(linha_up)

                if cnes and nome_estab and numero_apac:
                    break

            # ------------------------------------------------
            # Busca para baixo dentro do bloco da APAC
            # ------------------------------------------------

            k = i + 1

            while k < len(linhas):
                linha_down = limpar_linha(linhas[k])

                proxima_apac = extrair_apac(linha_down)

                if proxima_apac and proxima_apac != numero_apac:
                    break

                if "Processamento:" in linha_down:
                    processamento = extrair_processamento(linha_down)

                if "Producao:" in linha_down:
                    comp = extrair_producao(linha_down)

                if "Paciente" in linha_down and ":" in linha_down and "-" in linha_down:
                    cpf_tmp, cns_tmp, nome_tmp = extrair_paciente(linha_down)
                    cpf = cpf or cpf_tmp
                    cns = cns or cns_tmp
                    nome_pac = nome_pac or nome_tmp

                if "Responsavel:" in linha_down and "Dt.Nasc." in linha_down:
                    data_nasc = extrair_data_nasc(linha_down)

                # --------------------------------------------
                # Início da tabela de procedimentos
                # --------------------------------------------

                if re.search(
                    r"SQ\s+PROC\.\s+CBO\s+Qt\.Prz\.\s+Equipe\s+Cnes\s+Terc\.\s+Vl\.Prz\.\s+Qt\.Apvd\s+Vl\.Apvd\s+SITUACAO",
                    linha_down
                ):
                    k += 1
                    procedimentos_lidos_nesta_tabela = 0

                    while k < len(linhas):
                        linha_proc = limpar_linha(linhas[k])

                        # Fim da tabela
                        if re.match(r"^\s{22,}$", linha_proc):
                            break

                        # Ignora cabeçalhos espúrios
                        if (
                            "SISTEMA DE INFORMACOES AMBULATORIAIS" in linha_proc
                            or "SINTESE DA PRODUCAO DO APAC" in linha_proc
                        ):
                            k += 1
                            continue

                        # Nova APAC encontrada
                        proxima_apac_proc = extrair_apac(linha_proc)

                        if proxima_apac_proc and proxima_apac_proc != numero_apac:
                            k -= 1
                            break

                        # A linha de procedimento começa com SQ
                        seq_txt = extrair_faixa(linha_proc, 2, 3)

                        if not seq_txt.isdigit():
                            seq_txt = extrair_faixa(linha_proc, 1, 2)

                        if not seq_txt.isdigit():
                            break

                        seq = parse_numero_inteiro(seq_txt)

                        procedimento = normalizar_proc(
                            extrair_faixa(linha_proc, 5, 15)
                        )

                        cbo = extrair_faixa(linha_proc, 17, 22)
                        qt_prz_txt = extrair_faixa(linha_proc, 27, 30)
                        equipe = extrair_faixa(linha_proc, 32, 43)
                        cnes_terc = normalizar_cnes(
                            extrair_faixa(linha_proc, 45, 51)
                        )
                        vl_prz_txt = extrair_faixa(linha_proc, 57, 65)
                        qt_apvd_txt = extrair_faixa(linha_proc, 67, 72)
                        vl_apvd_txt = extrair_faixa(linha_proc, 74, 82)
                        situacao = extrair_faixa(linha_proc, 84, 102)

                        tp_proc = "PRINC" if seq == 1 else "SEC"

                        registro = {
                            "Processamento": processamento,
                            "CNES Estab. Exec.": cnes,
                            "Nome Estab. Exec.": nome_estab,
                            "Comp.": comp,
                            "N. APAC": numero_apac,
                            "CPF Pac.": cpf,
                            "CNS Pac.": cns,
                            "Nome Pac.": nome_pac,
                            "Data Nasc.": data_nasc,
                            "SEQ": seq,
                            "Procedimento": procedimento,
                            "Tp. Proc.": tp_proc,
                            "CBO": cbo,
                            "Qt.Prz.": parse_numero_inteiro(qt_prz_txt),
                            "Equipe": equipe,
                            "Cnes Terc.": cnes_terc,
                            "Nome Terc.": "",
                            "Vl.Prz.": parse_numero_real(vl_prz_txt),
                            "Qt.Apvd": parse_numero_inteiro(qt_apvd_txt),
                            "Vl.Apvd": parse_numero_real(vl_apvd_txt),
                            "SITUACAO": situacao
                        }

                        # Evita duplicidade de procedimento dentro da mesma APAC
                        chave = (
                            registro["N. APAC"],
                            registro["SEQ"],
                            registro["Procedimento"],
                            registro["CBO"],
                            registro["Qt.Apvd"],
                            registro["Vl.Apvd"],
                        )

                        if chave not in chaves_registro:
                            chaves_registro.add(chave)
                            dados.append(registro)
                            procedimentos_lidos_nesta_tabela += 1
                            apacs_oci_com_procedimentos.add(numero_apac)

                        k += 1

                    # Caso tenha achado cabeçalho vazio, continua procurando
                    if procedimentos_lidos_nesta_tabela == 0:
                        k += 1
                        continue

                k += 1

        i += 1

    df = pd.DataFrame(dados)

    if not df.empty:
        df = df.map(limpar_texto_excel)

    return (
        df,
        len(apacs_oci_identificadas),
        len(apacs_oci_com_procedimentos)
    )


# ============================================================
# COMPLEMENTOS APÓS EXTRAÇÃO
# ============================================================

def adicionar_nome_proced(df, mapa_nome_principal, mapa_nome_secundario):
    """
    Adiciona a coluna Nome_Proced após Procedimento.

    - PRINC usa mapa de procedimento principal;
    - SEC usa mapa de procedimento secundário.
    """
    if df.empty:
        return df

    df = df.copy()

    tipo_proc = df["Tp. Proc."].astype(str).str.upper().str.strip()
    procedimento_normalizado = df["Procedimento"].apply(normalizar_proc)

    df["Nome_Proced"] = ""

    df.loc[tipo_proc == "PRINC", "Nome_Proced"] = (
        procedimento_normalizado.map(mapa_nome_principal).fillna("")
    )

    df.loc[tipo_proc == "SEC", "Nome_Proced"] = (
        procedimento_normalizado.map(mapa_nome_secundario).fillna("")
    )

    colunas = df.columns.tolist()

    if "Nome_Proced" in colunas:
        colunas.remove("Nome_Proced")

        if "Procedimento" in colunas:
            indice = colunas.index("Procedimento") + 1
        else:
            indice = len(colunas)

        colunas.insert(indice, "Nome_Proced")
        df = df[colunas]

    return df


def organizar_coluna_data_nasc(df):
    """
    Move Data Nasc. para logo após Nome Pac.
    """
    if df.empty or "Data Nasc." not in df.columns or "Nome Pac." not in df.columns:
        return df

    df = df.copy()

    colunas = df.columns.tolist()
    colunas.remove("Data Nasc.")

    indice = colunas.index("Nome Pac.") + 1
    colunas.insert(indice, "Data Nasc.")

    return df[colunas]
# ============================================================
# PARTE 4/4
# MONTAGEM DAS ABAS, EXPORTAÇÃO E EXECUÇÃO PRINCIPAL
# ============================================================


# ============================================================
# ABA Vl_SIGTAP
# ============================================================

def montar_aba_vl_sigtap(df_dados_oci, principal_map, secundario_map):
    """
    Monta a aba Vl_SIGTAP.

    Para PRINC:
    - usa valor do procedimento principal.

    Para SEC:
    - usa valor do procedimento secundário compatível com o procedimento principal da APAC.
    """
    if df_dados_oci.empty:
        return pd.DataFrame()

    colunas_base = [
        "Processamento",
        "CNES Estab. Exec.",
        "Nome Estab. Exec.",
        "Equipe",
        "Comp.",
        "N. APAC",
        "Tp. Proc.",
        "Procedimento",
        "Nome_Proced",
        "Cnes Terc.",
        "Nome Terc.",
        "Qt.Apvd",
    ]

    df_vl = df_dados_oci[colunas_base].copy()

    df_vl["Vl. SIGTAP"] = 0.0
    df_vl["Total SIGTAP"] = 0.0

    proc_principal_por_apac = {}

    for idx, row in df_vl.iterrows():
        apac = str(row.get("N. APAC", "")).strip()
        tipo_proc = str(row.get("Tp. Proc.", "")).strip().upper()
        procedimento = normalizar_proc(row.get("Procedimento", ""))

        vl_sigtap = 0.0

        if tipo_proc == "PRINC":
            proc_principal_por_apac[apac] = procedimento
            vl_sigtap = float(principal_map.get(procedimento, 0.0))

        elif tipo_proc == "SEC":
            proc_principal = proc_principal_por_apac.get(apac, "")
            vl_sigtap = float(secundario_map.get((proc_principal, procedimento), 0.0))

        qt_apvd = pd.to_numeric(row.get("Qt.Apvd"), errors="coerce")
        qt_apvd = float(qt_apvd) if pd.notna(qt_apvd) else 0.0

        df_vl.at[idx, "Vl. SIGTAP"] = vl_sigtap
        df_vl.at[idx, "Total SIGTAP"] = qt_apvd * vl_sigtap

    return df_vl


# ============================================================
# ABA Vl_APAC_SIGTAP
# ============================================================

def montar_aba_vl_apac_sigtap(df_vl_sigtap):
    """
    Monta a aba Vl_APAC_SIGTAP.

    Agrupa os procedimentos secundários por APAC e soma o Total SIGTAP.
    Também recupera o valor OCI da linha PRINC da mesma APAC.
    """
    if df_vl_sigtap.empty:
        return pd.DataFrame(columns=[
            "Processamento",
            "Comp.",
            "CNES Estab. Exec.",
            "Nome Estab. Exec.",
            "Equipe",
            "N. APAC",
            "Total Vl. SIGTAP",
            "Vl. OCI",
        ])

    df_base = df_vl_sigtap.copy()
    df_base["Tp. Proc."] = df_base["Tp. Proc."].astype(str).str.upper().str.strip()

    df_sec = df_base[df_base["Tp. Proc."] == "SEC"].copy()
    df_princ = df_base[df_base["Tp. Proc."] == "PRINC"].copy()

    if df_sec.empty:
        return pd.DataFrame(columns=[
            "Processamento",
            "Comp.",
            "CNES Estab. Exec.",
            "Nome Estab. Exec.",
            "Equipe",
            "N. APAC",
            "Total Vl. SIGTAP",
            "Vl. OCI",
        ])

    df_sec["Total SIGTAP"] = pd.to_numeric(
        df_sec["Total SIGTAP"],
        errors="coerce"
    ).fillna(0.0)

    agrupado = (
        df_sec.groupby(
            [
                "Processamento",
                "Comp.",
                "CNES Estab. Exec.",
                "Nome Estab. Exec.",
                "Equipe",
                "N. APAC"
            ],
            as_index=False,
            dropna=False
        )["Total SIGTAP"].sum()
    )

    agrupado = agrupado.rename(columns={
        "Total SIGTAP": "Total Vl. SIGTAP"
    })

    df_princ["Total SIGTAP"] = pd.to_numeric(
        df_princ["Total SIGTAP"],
        errors="coerce"
    ).fillna(0.0)

    vl_oci_por_apac = (
        df_princ.groupby(
            "N. APAC",
            as_index=False,
            dropna=False
        )["Total SIGTAP"]
        .sum()
        .rename(columns={
            "Total SIGTAP": "Vl. OCI"
        })
    )

    agrupado = agrupado.merge(
        vl_oci_por_apac,
        on="N. APAC",
        how="left"
    )

    agrupado["Vl. OCI"] = pd.to_numeric(
        agrupado["Vl. OCI"],
        errors="coerce"
    ).fillna(0.0)

    return agrupado


# ============================================================
# PERCENTUAL NA ABA Vl_SIGTAP
# ============================================================

def aplicar_percentual_vl_sigtap(df_vl_sigtap, df_vl_apac_sigtap):
    """
    Calcula percentual dos procedimentos secundários.

    Percent. = Total SIGTAP da linha SEC / Total Vl. SIGTAP da APAC

    Mantém truncamento em duas casas, conforme regra anterior.
    """
    if df_vl_sigtap.empty:
        return df_vl_sigtap

    df_vl = df_vl_sigtap.copy()
    df_vl["Percent."] = ""

    mapa_total_apac = {}

    if not df_vl_apac_sigtap.empty:
        mapa_total_apac = (
            df_vl_apac_sigtap[["N. APAC", "Total Vl. SIGTAP"]]
            .drop_duplicates(subset=["N. APAC"])
            .set_index("N. APAC")["Total Vl. SIGTAP"]
            .to_dict()
        )

    tipo_proc = df_vl["Tp. Proc."].astype(str).str.upper().str.strip()
    mask_sec = tipo_proc == "SEC"

    total_apac_linha = df_vl["N. APAC"].map(mapa_total_apac)
    total_apac_linha = pd.to_numeric(total_apac_linha, errors="coerce").fillna(0.0)

    total_sigtap_linha = pd.to_numeric(
        df_vl["Total SIGTAP"],
        errors="coerce"
    ).fillna(0.0)

    percentual = (
        total_sigtap_linha / total_apac_linha.replace(0, pd.NA)
    ).fillna(0.0)

    percentual_truncado = percentual.apply(
        lambda x: math.floor(float(x) * 100) / 100 if pd.notna(x) else 0.0
    )

    df_vl.loc[mask_sec, "Percent."] = percentual_truncado[mask_sec]

    return df_vl


# ============================================================
# ABA Total_Terceiros
# ============================================================

def montar_aba_total_terceiros(df_dados_oci, df_vl_sigtap):
    """
    Monta a aba Total_Terceiros.

    Regras:
    - Cruza Total SIGTAP por APAC + Procedimento.
    - Cruza Percentual por APAC + Procedimento.
    - Para SEC:
        Vl. Proc. OCI = Total SIGTAP do PRINC da APAC * Percent.
    - Para PRINC:
        Vl. Proc. OCI = soma dos valores calculados das linhas SEC da APAC.
    """
    if df_dados_oci.empty:
        return pd.DataFrame()

    colunas_base = [
        "Processamento",
        "Comp.",
        "CNES Estab. Exec.",
        "Nome Estab. Exec.",
        "N. APAC",
        "Procedimento",
        "Nome_Proced",
        "Tp. Proc.",
        "Cnes Terc.",
        "Nome Terc.",
    ]

    df_total = df_dados_oci[colunas_base].copy()

    df_vl_join = df_vl_sigtap[
        [
            "N. APAC",
            "Procedimento",
            "Total SIGTAP"
        ]
    ].copy()

    df_total = df_total.merge(
        df_vl_join,
        on=[
            "N. APAC",
            "Procedimento"
        ],
        how="left"
    )

    df_total["Total SIGTAP"] = pd.to_numeric(
        df_total["Total SIGTAP"],
        errors="coerce"
    ).fillna(0.0)

    df_percent = df_vl_sigtap[
        [
            "N. APAC",
            "Procedimento",
            "Percent."
        ]
    ].copy()

    df_percent = df_percent.drop_duplicates(
        subset=[
            "N. APAC",
            "Procedimento"
        ]
    )

    df_total = df_total.merge(
        df_percent,
        on=[
            "N. APAC",
            "Procedimento"
        ],
        how="left"
    )

    df_total["Percent."] = pd.to_numeric(
        df_total["Percent."],
        errors="coerce"
    ).fillna(0.0).round(2)

    tipo_proc = df_total["Tp. Proc."].astype(str).str.upper().str.strip()

    total_princ_por_apac = (
        df_total[tipo_proc == "PRINC"]
        .groupby(
            "N. APAC",
            as_index=False,
            dropna=False
        )["Total SIGTAP"]
        .sum()
        .set_index("N. APAC")["Total SIGTAP"]
        .to_dict()
    )

    base_princ_linha = df_total["N. APAC"].map(total_princ_por_apac).fillna(0.0)

    vl_sec_linha = (
        base_princ_linha * df_total["Percent."]
    ).where(
        tipo_proc == "SEC",
        0.0
    )

    soma_sec_por_apac = vl_sec_linha.groupby(
        df_total["N. APAC"]
    ).sum()

    df_total["Vl. Proc. OCI"] = vl_sec_linha

    mask_princ = tipo_proc == "PRINC"

    df_total.loc[mask_princ, "Vl. Proc. OCI"] = (
        df_total.loc[mask_princ, "N. APAC"]
        .map(soma_sec_por_apac)
        .fillna(0.0)
    )

    df_total["Vl. Proc. OCI"] = pd.to_numeric(
        df_total["Vl. Proc. OCI"],
        errors="coerce"
    ).fillna(0.0).round(2)

    return df_total


# ============================================================
# FORMATAÇÃO FINAL DOS DATAFRAMES
# ============================================================

def formatar_datas_competencias(df):
    """
    Garante Processamento e Comp. no formato MM/YYYY.
    """
    if df.empty:
        return df

    df = df.copy()

    df = formatar_coluna_mm_aaaa(df, "Processamento")
    df = formatar_coluna_mm_aaaa(df, "Comp.")

    return df


def organizar_abas_finais(df_dados_oci, df_vl_sigtap, df_vl_apac_sigtap, df_total_terceiros):
    """
    Aplica limpeza final e padronização antes de salvar.
    """
    df_dados_oci = formatar_datas_competencias(df_dados_oci)
    df_vl_sigtap = formatar_datas_competencias(df_vl_sigtap)
    df_vl_apac_sigtap = formatar_datas_competencias(df_vl_apac_sigtap)
    df_total_terceiros = formatar_datas_competencias(df_total_terceiros)

    if not df_dados_oci.empty:
        df_dados_oci = df_dados_oci.map(limpar_texto_excel)

    if not df_vl_sigtap.empty:
        df_vl_sigtap = df_vl_sigtap.map(limpar_texto_excel)

    if not df_vl_apac_sigtap.empty:
        df_vl_apac_sigtap = df_vl_apac_sigtap.map(limpar_texto_excel)

    if not df_total_terceiros.empty:
        df_total_terceiros = df_total_terceiros.map(limpar_texto_excel)

    return df_dados_oci, df_vl_sigtap, df_vl_apac_sigtap, df_total_terceiros


# ============================================================
# SALVAR EXCEL
# ============================================================

def salvar_excel_final(df_dados_oci, df_vl_sigtap, df_vl_apac_sigtap, df_total_terceiros, caminho_saida):
    """
    Salva o arquivo Excel final com 4 abas.
    """
    (
        df_dados_oci,
        df_vl_sigtap,
        df_vl_apac_sigtap,
        df_total_terceiros
    ) = organizar_abas_finais(
        df_dados_oci,
        df_vl_sigtap,
        df_vl_apac_sigtap,
        df_total_terceiros
    )

    with pd.ExcelWriter(caminho_saida, engine="openpyxl") as writer:
        df_dados_oci.to_excel(
            writer,
            sheet_name="DADOS_OCI",
            index=False
        )

        df_vl_sigtap.to_excel(
            writer,
            sheet_name="Vl_SIGTAP",
            index=False
        )

        df_vl_apac_sigtap.to_excel(
            writer,
            sheet_name="Vl_APAC_SIGTAP",
            index=False
        )

        df_total_terceiros.to_excel(
            writer,
            sheet_name="Total_Terceiros",
            index=False
        )

        aplicar_destaques_excel(writer)


def aplicar_destaques_excel(writer):
    """
    Aplica destaque na aba Total_Terceiros quando:
    Tp. Proc. = PRINC
    e Total SIGTAP < Vl. Proc. OCI
    """
    if "Total_Terceiros" not in writer.book.sheetnames:
        return

    ws_total = writer.book["Total_Terceiros"]

    if ws_total.max_row < 2:
        return

    colunas_total = {
        cell.value: idx + 1
        for idx, cell in enumerate(ws_total[1])
    }

    col_tp_proc = colunas_total.get("Tp. Proc.")
    col_total_sigtap = colunas_total.get("Total SIGTAP")
    col_vl_proc_oci = colunas_total.get("Vl. Proc. OCI")

    if not col_tp_proc or not col_total_sigtap or not col_vl_proc_oci:
        return

    destaque = PatternFill(
        start_color="FFF2CC",
        end_color="FFF2CC",
        fill_type="solid"
    )

    for row_idx in range(2, ws_total.max_row + 1):
        tp_proc = ws_total.cell(row=row_idx, column=col_tp_proc).value
        total_sigtap = ws_total.cell(row=row_idx, column=col_total_sigtap).value
        vl_proc_oci = ws_total.cell(row=row_idx, column=col_vl_proc_oci).value

        tp_proc_txt = str(tp_proc or "").strip().upper()

        try:
            total_sigtap_num = float(total_sigtap or 0)
            vl_proc_oci_num = float(vl_proc_oci or 0)
        except:
            continue

        if tp_proc_txt == "PRINC" and total_sigtap_num < vl_proc_oci_num:
            ws_total.cell(
                row=row_idx,
                column=col_vl_proc_oci
            ).fill = destaque


# ============================================================
# EXECUÇÃO PRINCIPAL
# ============================================================

def main():
    # ============================================================
    # LOCALIZA AUTOMATICAMENTE A PASTA SIGTAP
    # ============================================================

    pasta_sigtap = os.path.join(
        pasta_do_programa(),
        "SIGTAP"
    )

    if not os.path.exists(pasta_sigtap):
        messagebox.showerror(
            "Pasta SIGTAP não encontrada",
            "Não foi encontrada uma pasta chamada SIGTAP no mesmo local do programa.\n\n"
            f"Local esperado:\n{pasta_sigtap}\n\n"
            "Crie uma pasta chamada SIGTAP no mesmo diretório do programa e coloque nela os arquivos:\n"
            "- tb_procedimento.txt\n"
            "- rl_procedimento_compativel.txt"
        )
        return
    

    # --------------------------------------------------------
    # Gera base SIGTAP OCI
    # --------------------------------------------------------

    df_sigtap, caminho_base_sigtap = gerar_base_oci_procedimentos_e_valores(
        pasta_sigtap
    )

    principal_map, secundario_map = carregar_mapas_valores_sigtap(
        df_sigtap
    )

    mapa_nome_principal, mapa_nome_secundario = carregar_mapas_nomes_procedimentos(
        df_sigtap
    )

    # --------------------------------------------------------
    # Seleciona arquivo SIA/APAC
    # --------------------------------------------------------

    messagebox.showinfo(
        "Selecionar arquivo SIA/APAC",
        "Selecione o arquivo SIA/APAC que será processado."
    )

    arquivo_sia = selecionar_arquivo(
        "Selecione o arquivo SIA/APAC"
    )

    if not arquivo_sia:
        messagebox.showwarning(
            "Cancelado",
            "Nenhum arquivo SIA/APAC foi selecionado."
        )
        return

    print("\n========================================")
    print(" PROCESSANDO ARQUIVO SIA/APAC")
    print("========================================")
    print(f"📄 Arquivo selecionado: {arquivo_sia}")

    (
        df_dados_oci,
        total_apacs_oci,
        total_apacs_com_procedimentos
    ) = parse_apac(arquivo_sia)

    if df_dados_oci.empty:
        messagebox.showwarning(
            "Nenhum dado extraído",
            "Nenhum dado OCI foi extraído do arquivo SIA/APAC selecionado."
        )
        return

    print(f"✅ APACs OCI identificadas: {total_apacs_oci}")
    print(f"✅ APACs OCI com procedimentos extraídos: {total_apacs_com_procedimentos}")
    print(f"✅ Linhas extraídas: {len(df_dados_oci)}")

    # --------------------------------------------------------
    # Complementa nomes de estabelecimentos
    # --------------------------------------------------------

    mapa_cnes = carregar_mapa_estabelecimentos()
    df_dados_oci = preencher_nomes_estabelecimentos(
        df_dados_oci,
        mapa_cnes
    )

    # --------------------------------------------------------
    # Complementa nomes de procedimentos
    # --------------------------------------------------------

    df_dados_oci = adicionar_nome_proced(
        df_dados_oci,
        mapa_nome_principal,
        mapa_nome_secundario
    )

    df_dados_oci = organizar_coluna_data_nasc(
        df_dados_oci
    )

    # --------------------------------------------------------
    # Monta abas calculadas
    # --------------------------------------------------------

    print("\n========================================")
    print(" MONTANDO ABAS DO EXCEL")
    print("========================================")

    df_vl_sigtap = montar_aba_vl_sigtap(
        df_dados_oci,
        principal_map,
        secundario_map
    )

    df_vl_apac_sigtap = montar_aba_vl_apac_sigtap(
        df_vl_sigtap
    )

    df_vl_sigtap = aplicar_percentual_vl_sigtap(
        df_vl_sigtap,
        df_vl_apac_sigtap
    )

    df_total_terceiros = montar_aba_total_terceiros(
        df_dados_oci,
        df_vl_sigtap
    )

    print(f"✅ Aba DADOS_OCI: {len(df_dados_oci)} linhas")
    print(f"✅ Aba Vl_SIGTAP: {len(df_vl_sigtap)} linhas")
    print(f"✅ Aba Vl_APAC_SIGTAP: {len(df_vl_apac_sigtap)} linhas")
    print(f"✅ Aba Total_Terceiros: {len(df_total_terceiros)} linhas")

    # --------------------------------------------------------
    # Salvar Excel final
    # --------------------------------------------------------

    caminho_saida = salvar_arquivo_excel(
        "Salvar arquivo final",
        "OCI_Producao_Prestadores.xlsx"
    )

    if not caminho_saida:
        messagebox.showwarning(
            "Cancelado",
            "O arquivo final não foi salvo."
        )
        return

    salvar_excel_final(
        df_dados_oci,
        df_vl_sigtap,
        df_vl_apac_sigtap,
        df_total_terceiros,
        caminho_saida
    )

    print("\n========================================")
    print(" PROCESSAMENTO CONCLUÍDO")
    print("========================================")
    print(f"✅ Arquivo final gerado: {caminho_saida}")
    print(f"📄 Base SIGTAP gerada: {caminho_base_sigtap}")
    print(f"✅ Total de APACs OCI identificadas: {total_apacs_oci}")
    print(f"✅ Total de APACs OCI com procedimentos no Excel: {total_apacs_com_procedimentos}")

    messagebox.showinfo(
        "Processamento concluído",
        "Arquivo gerado com sucesso!\n\n"
        f"Arquivo final:\n{caminho_saida}\n\n"
        f"Base SIGTAP gerada:\n{caminho_base_sigtap}\n\n"
        f"Total de APACs OCI identificadas: {total_apacs_oci}\n"
        f"Total de APACs OCI com procedimentos no Excel: {total_apacs_com_procedimentos}"
    )


if __name__ == "__main__":
    main()